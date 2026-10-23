"""Dados FICTÍCIOS do painel de rádio, gerados na hora (modo demonstração).

Gera as três planilhas no formato do EasyMedia (histórico das emissoras próprias desde jan–mar/2020, ranking de
concorrência e totais de mercado desde 2024) e um parametros.json com números "oficiais" e previsão do BI fictícios.
As concorrentes têm nomes inventados. Semente fixa: os mesmos números a cada geração."""
import json
import os

import numpy as np
import openpyxl
import pandas as pd

MESES_UP = "JAN FEV MAR ABR MAI JUN JUL AGO SET OUT NOV DEZ".split()
AVISO = "DADOS FICTÍCIOS · gerados para demonstração · não representam a audiência real"


def banco_txt(fim):
    ini = fim - pd.DateOffset(months=2)
    return f"SAL - {MESES_UP[ini.month - 1]}/{ini.year} A {MESES_UP[fim.month - 1]}/{fim.year}"


def serie_mensal(rng, meses, nivel, tend_ano, amp_saz, fase, ruido, ar=.55):
    n = len(meses)
    e = np.zeros(n)
    for i in range(n):
        e[i] = (e[i - 1] * ar if i else 0) + rng.normal(0, ruido)
    t = np.arange(n) / 12
    saz = amp_saz * np.sin(2 * np.pi * (np.array([m.month for m in meses]) - fase) / 12)
    return nivel * (1 + tend_ano * t) * (1 + saz) * (1 + e)


def gera(destino="."):
    """Cria destino/dados/*.xlsx e destino/parametros.json."""
    RNG = np.random.default_rng(20261021)
    pasta = os.path.join(destino, "dados")
    os.makedirs(pasta, exist_ok=True)
    meses = pd.date_range("2019-11-01", "2026-09-01", freq="MS")      # 1º banco: jan–mar/2020
    proprias = {"BAHIA FM/WEB": (24500, -.028, .045, 3), "GFM/WEB": (35500, -.022, .05, 4)}
    concorrentes = {  # nomes inventados
        "ATLÂNTICA FM/WEB": (31500, .004, .04, 5), "FAROL FM/WEB": (17500, -.03, .035, 2),
        "ONDA SUL FM/WEB": (11800, .015, .05, 6), "BRISA FM/WEB": (9300, -.01, .04, 3),
        "MARÉ FM/WEB": (7600, -.005, .045, 4), "HORIZONTE FM/WEB": (7100, .02, .04, 7),
        "CIDADE ALTA FM/WEB": (5400, -.02, .05, 5), "LITORAL FM/WEB": (5200, .01, .04, 2),
        "AURORA FM/WEB": (5000, -.015, .035, 8), "CORAL FM/WEB": (4400, .03, .05, 3),
        "SERTÃO FM/WEB": (3700, -.01, .04, 6), "PRIMAVERA AM/WEB": (2100, -.04, .04, 4),
        "VENTO NORTE FM/WEB": (1300, .02, .06, 5), "PONTE FM/WEB": (700, .0, .08, 7)}
    todas = {**proprias, **concorrentes}
    mensal = {k: serie_mensal(RNG, meses, *v, ruido=.035) for k, v in todas.items()}
    fins = [f for f in pd.date_range("2020-03-01", "2026-09-01", freq="MS")]

    def opm(nome, fim):
        i = meses.get_loc(fim)
        v = mensal[nome][i - 2:i + 1].mean() * (1 + RNG.normal(0, .015))
        return float(round(v))

    bancos = {k: {f: opm(k, f) for f in fins} for k in todas}

    # ---- histórico das duas emissoras próprias (uma aba por ano, pelo mês inicial do banco)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for ano in range(2020, 2027):
        ws = wb.create_sheet(str(ano))
        cab = [["Rádio Recall"], [AVISO], ["EasyMedia4 (DEMO)"], [], [], [None, None, "GRANDE SALVADOR"], [None, None, "05-05"],
               [None, None, "TODOS OS DIAS"], [None, None, "OPM#"], ["Agrupamento", "Período", "SAL - SEXO AMBOS"]]
        for r in cab:
            ws.append(r)
        for nome in proprias:
            fs = [f for f in fins if (f - pd.DateOffset(months=2)).year == ano]
            for k, f in enumerate(fs):
                ws.append([f"SAL - {nome}" if k == 0 else None, banco_txt(f), bancos[nome][f]])
        ws.append([])
        ws.append([f"Fonte: {AVISO}"])
    wb.save(os.path.join(pasta, "DEMO_Historico_Radio_2020_2026.xlsx"))

    # ---- ranking de concorrência (2024 em diante, sem nov–jan e dez–fev, como no EasyMedia)
    fins_r = [f for f in fins if f >= pd.Timestamp("2024-03-01") and f.month not in (1, 2)]
    total_mercado = {f: sum(bancos[k][f] for k in todas) * 1.04 for f in fins_r}
    alc_fator = {k: RNG.uniform(17, 25) for k in todas}
    tmed_base = {k: RNG.uniform(105, 205) for k in todas}

    def metricas(nome, f):
        o = bancos[nome][f]
        if nome == "PONTE FM/WEB" and RNG.random() < .25:
            return ["NA"] * 4
        sht = o / total_mercado[f] * 100
        alc = o * alc_fator[nome] * (1 + RNG.normal(0, .02))
        mins = tmed_base[nome] * (1 + RNG.normal(0, .03))
        tm = f"{int(mins // 60):02d}:{int(mins % 60):02d}:{int(RNG.integers(0, 60)):02d}"
        baixa = o < 1500
        p = "#" if baixa else ""
        return [f"{p}{int(o)}" if baixa else int(o), f"{p}{sht:.2f}".replace(".", ",") if baixa else round(sht, 2),
                f"{p}{alc:.2f}".replace(".", ",") if baixa else round(alc, 2), f"{p}{tm}" if baixa else tm]

    def planilha_larga(linhas_nomes, arquivo, valores):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "RANKING GERAL - 05-05_05H00"
        for r in [["Rádio Recall"], [AVISO], ["EasyMedia4 (DEMO)"], [], [], [None, "GRANDE SALVADOR"], [None, "05-05"],
                  [None, "TODOS OS DIAS"], [None, "SAL - SEXO AMBOS"]]:
            ws.append(r)
        per = [None]
        met = ["Agrupamento"]
        for f in fins_r:
            per += [banco_txt(f), None, None, None]
            met += ["OPM#", "SHT%", "ALCM#", "TMED#"]
        ws.append(per)
        ws.append(met)
        for nome in linhas_nomes:
            linha = [f"SAL - {nome}"]
            for f in fins_r:
                linha += valores(nome, f)
            ws.append(linha)
        ws.append([])
        ws.append([f"Fonte: {AVISO}"])
        wb.save(os.path.join(pasta, arquivo))

    planilha_larga(sorted(todas), "DEMO_Concorrencia_2024_2026.xlsx", metricas)

    # ---- totais de mercado (consulta)
    def totais(nome, f):
        soma = sum(bancos[k][f] for k in todas)
        fator = {"TOTAL EMISSORAS": 1.0, "TOTAL FM": .9, "TOTAL WEB": .075, "TOTAL RADIO": 1.04}[nome]
        o = soma * fator * (1 + RNG.normal(0, .006))
        sht = o / total_mercado[f] * 100
        mins = 190 + RNG.normal(0, 4)
        return [int(round(o)), round(min(sht, 99.5), 2), round(o * 11 * (1 + RNG.normal(0, .01)), 2),
                f"{int(mins // 60):02d}:{int(mins % 60):02d}:{int(RNG.integers(0, 60)):02d}"]

    planilha_larga(["TOTAL EMISSORAS", "TOTAL FM", "TOTAL WEB", "TOTAL RADIO"], "DEMO_Totais_2024_2026.xlsx", totais)

    # ---- parâmetros: números "oficiais" e previsão do BI (fictícios, coerentes com os bancos acima)
    oficial = {}
    for nome, chave in (("GFM/WEB", "GFM"), ("BAHIA FM/WEB", "Bahia FM")):
        o = {}
        for ano in (2023, 2024, 2025):
            internos = [bancos[nome][pd.Timestamp(ano, m, 1)] for m in range(3, 13)]
            o[str(ano)] = int(round(np.mean(internos)))
            o[f"1S{str(ano)[2:]}"] = int(round(np.mean([bancos[nome][pd.Timestamp(ano, m, 1)] for m in (3, 4, 5, 6)])))
            o[f"2S{str(ano)[2:]}"] = int(round(np.mean([bancos[nome][pd.Timestamp(ano, m, 1)] for m in (9, 10, 11, 12)])))
        oficial[chave] = o
    par = {"demo": True, "oficial": oficial,
           "previsao_bi": {"Bahia FM": {"S1": round(oficial["Bahia FM"]["2025"] * .9, -2), "S2": round(oficial["Bahia FM"]["2025"] * .92, -2),
                                        "ano": round(oficial["Bahia FM"]["2025"] * .9, -2)},
                           "GFM": {"S1": round(oficial["GFM"]["2025"] * .97, -2), "S2": round(oficial["GFM"]["2025"] * .95, -2),
                                   "ano": round(oficial["GFM"]["2025"] * .96, -2)}},
           "dados": {"descricao": "DADOS FICTÍCIOS de demonstração", "atualizado_em": "21/10/2026"}}
    json.dump(par, open(os.path.join(destino, "parametros.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    return par


if __name__ == "__main__":
    gera(".")
    print("dados fictícios gerados em ./dados e ./parametros.json")
